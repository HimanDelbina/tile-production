from io import BytesIO
from decimal import Decimal
from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import BarChart,Reference
from openpyxl.utils import get_column_letter
from .dates import jalali

METHOD='سهم از نتایج فیلترشده = متراژ گروه / جمع نتایج × ۱۰۰. ترکیب درجات با حذف فیلتر درجه و حفظ سایر فیلترها محاسبه می‌شود. سهم داخلی کارخانه نسبت به جمع همان کارخانه است. درصدها هنگام نمایش گرد می‌شوند؛ جمع ممکن است اندکی با ۱۰۰٪ اختلاف داشته باشد. — یعنی داده کافی نیست.'

def metadata(ctx):
    factories = '، '.join(f['label'] for f in ctx['summary']['factories']) or 'بدون تولید در محدوده مجاز'
    filters = ' | '.join(f['label'] + ': ' + f['value'] for f in ctx['active_filters']) or 'همه تاریخ‌ها و همه کارخانه‌های مجاز'
    return ['گزارش مدیریتی تولید کاشی', 'کارخانه‌های محدوده گزارش: ' + factories, 'فیلترها: ' + filters, 'واحد: مترمربع | زمان تهیه: ' + jalali(timezone.localdate()) + ' ' + timezone.localtime().strftime('%H:%M'), METHOD]

def excel(ctx):
    wb=Workbook();ws=wb.active;ws.title='گزارش تولید';ws.sheet_view.rightToLeft=True
    table=ctx['table'];width=max(4,len(table['headers']))
    for i,text in enumerate(metadata(ctx),1):
        ws.cell(i,1,text);ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=width);ws.row_dimensions[i].height=32 if i!=5 else 55
    ws.cell(6,1,'کل نتایج فیلترشده (مترمربع)');ws.cell(6,2,float(ctx['summary']['total']));ws.cell(6,2).number_format='#,##0.00'
    start=8
    for j,h in enumerate(table['headers'],1): ws.cell(start,j,h)
    for i,row in enumerate(table['body'],start+1):
        for j,value in enumerate(row['cells'],1):
            if j-1 in table['percent_cols']:
                value=float(value/100) if value is not None else None
            elif isinstance(value,Decimal): value=float(value)
            cell=ws.cell(i,j,value if value is not None else '—')
            # Untrusted user text must never become an Excel formula.
            if isinstance(value,str): cell.data_type='s'
            if j-1 in table['percent_cols']: cell.number_format='0.00%'
            elif j-1 in table['area_cols']: cell.number_format='#,##0.00'
    ws.freeze_panes='D9';ws.auto_filter.ref=f'A8:{get_column_letter(len(table["headers"]))}{max(8,ws.max_row)}'
    ws.print_title_rows='1:8';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    summary=wb.create_sheet('خلاصه مدیریتی');summary.sheet_view.rightToLeft=True
    summary.append(['شاخص','مقدار','مبنا']);summary.append(['کل تولید',float(ctx['summary']['total']),'مترمربع؛ نتایج فیلترشده']);summary.append(['سهم درجه یک',float(ctx['summary']['first_share']/100) if ctx['summary']['first_share'] is not None else None,'همه درجات؛ سایر فیلترها حفظ می‌شوند']);summary['B3'].number_format='0.00%'
    chart_ranges=[]
    for title,key in [('کارخانه‌ها','factories'),('سایزها','sizes'),('درجات','grades')]:
        summary.append([]);s=summary.max_row+1;summary.append([title,'متراژ (مترمربع)','سهم از نتایج فیلترشده'])
        for item in ctx['summary'][key]:
            summary.append([item['label'],float(item['area']),float(item['share']/100) if item['share'] is not None else None]);summary.cell(summary.max_row,2).number_format='#,##0.00';summary.cell(summary.max_row,3).number_format='0.00%'
        if summary.max_row>s: chart_ranges.append((title,s,summary.max_row))
    for n,(title,s,end) in enumerate(chart_ranges):
        chart=BarChart();chart.title=title;chart.y_axis.title='مترمربع';chart.add_data(Reference(summary,min_col=2,min_row=s,max_row=end),titles_from_data=True);chart.set_categories(Reference(summary,min_col=1,min_row=s+1,max_row=end));chart.width=18;chart.height=9;summary.add_chart(chart,f'E{2+n*19}')
    for sheet in wb:
        for row in sheet:
            for c in row:
                if isinstance(c.value,str): c.data_type='s'
                c.font=Font(name='Vazirmatn',size=11,color='243C53');c.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F2F7FA')
        headrow=start if sheet==ws else 1
        for c in sheet[headrow]: c.fill=PatternFill('solid',fgColor='122C48');c.font=Font(name='Vazirmatn',size=11,bold=True,color='FFFFFF')
        for i in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(i)].width=24
        sheet.column_dimensions['A'].width=28
    out=BytesIO();wb.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']='attachment; filename="tile-production.xlsx"';return response

def print_context():
    return {'regular_font':(settings.BASE_DIR/'static/vendor/Vazirmatn-Regular.ttf').as_uri(),'bold_font':(settings.BASE_DIR/'static/vendor/Vazirmatn-Bold.ttf').as_uri()}

def prepare_print(ctx):
    table=ctx['table'];n=len(table['headers'])
    if table['group']=='matrix' and n>11:
        bands=[]
        for start in range(3,n-2,6):
            indices=[0,1,2]+list(range(start,min(start+6,n-2)))+[n-2,n-1]
            bands.append({'headers':[table['headers'][i] for i in indices],'rows':[[r['cells'][i] for i in indices] for r in table['body']]})
    else: bands=[{'headers':table['headers'],'rows':[r['cells'] for r in table['body']]}]
    ctx.update(meta=metadata(ctx),bands=bands,method=METHOD)
    return ctx

def pdf(ctx):
    import os, sys
    if sys.platform == 'win32':
        for p in [
            r"C:\Program Files\GTK3-Runtime Win64\bin",
            r"C:\GTK3\bin",
            r"C:\msys64\ucrt64\bin",
            r"C:\msys64\mingw64\bin",
        ]:
            if os.path.isdir(p):
                try:
                    os.add_dll_directory(p)
                except Exception:
                    pass

    from weasyprint import HTML,default_url_fetcher
    ctx=prepare_print(ctx)
    # No network or arbitrary local files may be fetched by the PDF renderer.
    allowed={ctx['regular_font'],ctx['bold_font']}
    def fetch(url,*args,**kwargs):
        if url not in allowed: raise ValueError('منبع PDF مجاز نیست')
        return default_url_fetcher(url,*args,**kwargs)
    data=HTML(string=render_to_string('production/print.html',ctx),base_url=str(settings.BASE_DIR),url_fetcher=fetch).write_pdf()
    response=HttpResponse(data,content_type='application/pdf');response['Content-Disposition']='attachment; filename="tile-production.pdf"';return response
