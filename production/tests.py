import uuid
import os
from decimal import Decimal as D
from datetime import date, timedelta
from io import BytesIO
from unittest import skipUnless

from django.test import TestCase, Client, TransactionTestCase
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError, PermissionDenied
from django.http import QueryDict
from django.db import connection, close_old_connections
from django.urls import reverse

from .models import Factory, Size, Grade, Profile, Production, Audit, Submission
from .forms import ProductionForm, FilterForm
from .dates import parse_jalali, period, jalali
from .services import create_batch, change_record, duplicate_rows
from .reporting import summarize, report_table, filtered


class ProductionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser('admin', 'admin@example.test', 'Strong-password-923!')
        cls.f1 = Factory.objects.create(name='کارخانه اول')
        cls.f2 = Factory.objects.create(name='کارخانه دوم')
        cls.s = Size.objects.create(width=60, length=120)
        cls.g1 = Grade.objects.create(name='درجه ۱', rank=1, color='#159b9a')
        cls.g2 = Grade.objects.create(name='درجه ۲', rank=2, color='#eba743')
        cls.day = parse_jalali('۱۴۰۵/۰۶/۱۸')
        # Create initial productions without design
        for f, areas in [(cls.f1, [800, 200]), (cls.f2, [300, 200])]:
            rows = [{'size': cls.s, 'grade': g, 'area': D(area), 'notes': ''}
                    for g, area in zip([cls.g1, cls.g2], areas)]
            create_batch(cls.admin, f, cls.day, rows, uuid.uuid4())
        cls.entry = get_user_model().objects.create_user('entry', password='Strong-password-923!')
        p = Profile.objects.create(user=cls.entry, role='entry')
        p.factories.add(cls.f1)
        cls.viewer = get_user_model().objects.create_user('viewer', password='Strong-password-923!')
        p = Profile.objects.create(user=cls.viewer, role='viewer')
        p.factories.add(cls.f1)

    def setUp(self):
        self.client.force_login(self.admin)

    def row(self, area='1.25'):
        return {'size': self.s, 'grade': self.g1, 'area': D(area), 'notes': ''}

    def summary(self, data=None, user=None):
        return summarize(user or self.admin, data or {}, QueryDict())

    def payload(self, factory=None, token=None, area='۱۲٫۳۴'):
        return {
            'factory': (factory or self.f1).pk,
            'date': '۱۴۰۵/۰۶/۱۸',
            'token': str(token or uuid.uuid4()),
            'rows-TOTAL_FORMS': '1',
            'rows-INITIAL_FORMS': '0',
            'rows-MIN_NUM_FORMS': '1',
            'rows-MAX_NUM_FORMS': '100',
            'rows-0-size': self.s.pk,
            'rows-0-grade': self.g1.pk,
            'rows-0-area': area,
            'rows-0-notes': '',
        }

    def test_required_scenario(self):
        s = self.summary()
        self.assertEqual(s['total'], D(1500))
        self.assertEqual(s['first_share'].quantize(D('.01')), D('73.33'))
        f = {x['id']: x for x in s['factories']}
        self.assertEqual(f[self.f1.pk]['area'], D(1000))
        self.assertEqual(f[self.f2.pk]['area'], D(500))
        self.assertEqual(f[self.f1.pk]['share'].quantize(D('.01')), D('66.67'))
        self.assertEqual(self.summary({'factory': [self.f1]})['first_share'], D(80))
        self.assertEqual(self.summary({'factory': [self.f2]})['first_share'], D(60))

    def test_grade_filter_denominator(self):
        s = self.summary({'grade': [self.g1]})
        self.assertEqual(s['total'], D(1100))
        self.assertEqual(s['base_total'], D(1500))
        self.assertEqual(s['first_share'].quantize(D('.01')), D('73.33'))
        s = self.summary({'grade': [self.g2]})
        self.assertEqual(s['total'], D(400))
        self.assertEqual(s['first_share'].quantize(D('.01')), D('73.33'))

    def test_composite_filters(self):
        Production.objects.filter(factory=self.f1, area=800).update(notes='W-100')
        d = {
            'factory': [self.f1],
            'size': [self.s],
            'grade': [self.g1],
            'start': self.day,
            'end': self.day,
            'minimum': D(750),
            'maximum': D(850),
            'author': self.admin,
            'q': 'W-100',
        }
        self.assertEqual(self.summary(d)['total'], 800)
        d['q'] = 'unknown'
        self.assertEqual(self.summary(d)['total'], 0)

    def test_empty_denominator(self):
        s = self.summary({'start': self.day + timedelta(days=1)})
        self.assertEqual(s['total'], 0)
        self.assertIsNone(s['first_share'])
        self.assertEqual(len(s['factories']), 2)

    def test_jalali_dates(self):
        self.assertEqual(parse_jalali('۱۴۰۵/۰۶/۱۸'), date(2026, 9, 9))
        self.assertEqual(parse_jalali('١٤٠٥/٠٦/١٨'), self.day)
        self.assertEqual(jalali(parse_jalali('1403/12/30')), '1403/12/30')
        for value in ['1404/12/30', '1405/07/31', '1405/13/01', '2026-09-09', 'not-a-date']:
            with self.assertRaises(ValidationError):
                parse_jalali(value)

    def test_jalali_boundaries(self):
        a, b = period('lastmonth', parse_jalali('1404/01/01'))
        self.assertEqual((jalali(a), jalali(b)), ('1403/12/01', '1403/12/30'))
        a, b = period('month', parse_jalali('1405/06/18'))
        self.assertEqual((jalali(a), jalali(b)), ('1405/06/01', '1405/06/31'))
        a, b = period('year', self.day)
        self.assertEqual((jalali(a), jalali(b)), ('1405/01/01', '1405/12/29'))
        a, b = period('week', self.day)
        self.assertEqual(a.weekday(), 5)

    def test_month_groups_use_jalali(self):
        create_batch(self.admin, self.f1, parse_jalali('1405/07/01'), [self.row()], uuid.uuid4())
        t = report_table(self.summary(), {'group': 'month'})
        self.assertEqual({r['cells'][0] for r in t['body']}, {'1405/06', '1405/07'})

    def test_decimal_validation(self):
        for value, valid in [('۱۲٫۳۴', True), ('12.345', False), ('-1', False), ('0', False), ('NaN', False), ('Infinity', False)]:
            form = ProductionForm({'size': self.s.pk, 'grade': self.g1.pk, 'area': value})
            self.assertEqual(form.is_valid(), valid, form.errors)

    def test_atomic_rollback(self):
        count = Production.objects.count()
        key = uuid.uuid4()
        with self.assertRaises(ValidationError):
            create_batch(self.admin, self.f1, self.day, [self.row(), self.row('-1')], key)
        self.assertEqual(Production.objects.count(), count)
        self.assertFalse(Submission.objects.filter(key=key).exists())

    def test_idempotency(self):
        key = uuid.uuid4()
        count = Production.objects.count()
        self.assertTrue(create_batch(self.admin, self.f1, self.day, [self.row()], key))
        self.assertFalse(create_batch(self.admin, self.f1, self.day, [self.row()], key))
        self.assertEqual(Production.objects.count(), count + 1)
        with self.assertRaises(ValidationError):
            create_batch(self.admin, self.f1, self.day, [self.row('2')], key)

    def test_multiple_runs_allowed_and_duplicate_warning(self):
        data = self.row('800')
        self.assertEqual(duplicate_rows(self.admin, self.f1, self.day, [data]), [1])
        create_batch(self.admin, self.f1, self.day, [data], uuid.uuid4())
        self.assertEqual(self.summary()['total'], 2300)

    def test_post_replay_and_warning(self):
        p = self.payload()
        count = Production.objects.count()
        self.assertEqual(self.client.post('/production/new/', p).status_code, 302)
        self.assertEqual(self.client.post('/production/new/', p).status_code, 302)
        self.assertEqual(Production.objects.count(), count + 1)
        p['token'] = str(uuid.uuid4())
        response = self.client.post('/production/new/', p)
        self.assertContains(response, 'مشابه')
        self.assertEqual(Production.objects.count(), count + 1)
        p['confirm_duplicate'] = 'on'
        self.assertEqual(self.client.post('/production/new/', p).status_code, 302)
        self.assertEqual(Production.objects.count(), count + 2)

    def test_invalid_batch_form_stores_nothing(self):
        p = self.payload()
        p.update({
            'rows-TOTAL_FORMS': '2',
            'rows-1-size': self.s.pk,
            'rows-1-grade': self.g1.pk,
            'rows-1-area': '0',
        })
        n = Production.objects.count()
        self.assertEqual(self.client.post('/production/new/', p).status_code, 200)
        self.assertEqual(Production.objects.count(), n)

    def test_factory_scope(self):
        self.assertEqual(self.summary(user=self.viewer)['total'], 1000)
        obj = Production.objects.filter(factory=self.f2).first()
        self.client.force_login(self.entry)
        for route in ['edit', 'delete', 'history']:
            self.assertEqual(self.client.get(reverse(route, args=[obj.pk])).status_code, 404)
        self.assertEqual(self.client.post('/production/new/', self.payload(self.f2)).status_code, 200)
        self.assertEqual(self.client.get('/reports/', {'factory': self.f2.pk}).status_code, 400)
        self.assertEqual(self.client.get('/export/xlsx/', {'factory': self.f2.pk}).status_code, 400)
        with self.assertRaises(PermissionDenied):
            create_batch(self.entry, self.f2, self.day, [self.row()], uuid.uuid4())

    def test_viewer_no_writes(self):
        self.client.force_login(self.viewer)
        obj = Production.objects.filter(factory=self.f1).first()
        urls = ['/production/new/', reverse('edit', args=[obj.pk]), reverse('delete', args=[obj.pk]), '/users/']
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 403)

    def test_soft_delete_and_audit(self):
        obj = Production.objects.filter(area=800).get()
        change_record(self.admin, obj.pk, version=1, delete=True)
        self.assertEqual(self.summary()['total'], 700)
        self.assertTrue(Production.objects.filter(pk=obj.pk).exists())
        self.assertEqual(Audit.objects.filter(production=obj, action='delete').count(), 1)
        self.assertEqual(self.client.get(reverse('history', args=[obj.pk])).status_code, 200)

    def test_optimistic_edit_conflict(self):
        obj = Production.objects.filter(factory=self.f1).first()
        change_record(self.admin, obj.pk, {'area': D('15.25')}, version=1)
        with self.assertRaises(ValidationError):
            change_record(self.admin, obj.pk, {'area': D('99')}, version=1)
        obj.refresh_from_db()
        self.assertEqual(obj.area, D('15.25'))
        self.assertEqual(obj.version, 2)

    def test_matrix_dynamic_grades(self):
        Grade.objects.create(name='درجه ویژه', rank=8, color='#445566')
        t = report_table(self.summary(), {'group': 'matrix'})
        self.assertEqual(len(t['headers']), 10)
        self.assertTrue(any('ویژه' in h for h in t['headers']))
        s = self.summary({'grade': [self.g1]})
        t = report_table(s, {'group': 'matrix'})
        r = next(r for r in t['body'] if r['cells'][0] == self.f1.name)
        self.assertEqual(r['cells'][3], 80)
        self.assertEqual(r['cells'][5], 20)

    def test_report_routes_and_groups(self):
        urls = ['/', '/reports/', '/production/new/', '/production/new/?mode=single', '/master/factories/', '/users/', '/password/']
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        for group in ['day', 'month', 'year', 'factory', 'size', 'grade', 'type', 'matrix']:
            self.assertEqual(self.client.get('/reports/', {'group': group}).status_code, 200, group)

    def test_invalid_filter_rejected(self):
        queries = [
            {'start': '1405/13/01'},
            {'start': '1405/06/19', 'end': '1405/06/18'},
            {'minimum': '100', 'maximum': '50'},
            {'grade': '9999'},
        ]
        for q in queries:
            self.assertEqual(self.client.get('/reports/', q).status_code, 400)

    def test_calendar_api(self):
        r = self.client.get('/calendar/', {'year': '۱۴۰۳', 'month': '۱۲'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['days'], 30)
        self.assertEqual(self.client.get('/calendar/', {'year': '1405', 'month': '13'}).status_code, 400)

    def test_csrf(self):
        c = Client(enforce_csrf_checks=True)
        c.force_login(self.admin)
        self.assertEqual(c.post('/production/new/', self.payload()).status_code, 403)

    def test_auth_redirect(self):
        self.client.logout()
        for url in ['/', '/reports/', '/export/pdf/', '/calendar/']:
            self.assertEqual(self.client.get(url).status_code, 302)

    def test_exports_all_rows_and_numbers(self):
        from openpyxl import load_workbook
        rows = [self.row(str(i) + '.25') for i in range(1, 31)]
        create_batch(self.admin, self.f1, self.day, rows, uuid.uuid4())
        response = self.client.get('/reports/', {'page_size': 25})
        self.assertEqual(len(response.context['page']), 25)
        self.assertEqual(response.context['summary']['total'], D('1972.50'))
        response = self.client.get('/export/xlsx/', {'page_size': 25})
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.active.max_row, 42)
        self.assertEqual(wb.active['B6'].value, 1972.5)
        self.assertEqual(wb.active['E9'].data_type, 'n')
        response = self.client.get('/export/xlsx/', {'group': 'factory'})
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.active['C9'].number_format, '0.00%')

    def test_excel_formula_injection(self):
        create_batch(self.admin, self.f1, self.day, [{**self.row(), 'notes': '=HYPERLINK("https://example.test")'}], uuid.uuid4())
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(self.client.get('/export/xlsx/').content))
        self.assertEqual(wb.active['G9'].data_type, 's')

    def test_print_and_pdf(self):
        response = self.client.get('/export/print/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '۱٬۵۰۰٫۰۰')
        response = self.client.get('/export/pdf/')
        self.assertEqual(response.status_code, 200)
        if response['Content-Type'] == 'application/pdf':
            self.assertTrue(response.content.startswith(b'%PDF'))
        else:
            self.assertContains(response, 'چاپ گزارش')

    def test_export_scope(self):
        self.client.force_login(self.viewer)
        r = self.client.get('/export/print/')
        self.assertContains(r, 'کارخانه اول')
        self.assertNotContains(r, 'کارخانه دوم')

    def test_month_shortcut_normalizes_query(self):
        r = self.client.get('/reports/', {'preset': 'lastmonth'})
        self.assertNotIn('preset', r.context['query'])
        self.assertIn('start=', r.context['query'])

    def test_update_and_users_views(self):
        obj = Production.objects.filter(factory=self.f1).first()
        r = self.client.get(reverse('edit', args=[obj.pk]))
        self.assertEqual(r.status_code, 200)
        p = {
            'factory': self.f1.pk,
            'date': '1405/06/18',
            'token': uuid.uuid4(),
            'version': 1,
            'size': self.s.pk,
            'grade': self.g1.pk,
            'area': '۵۵٫۲۵',
            'notes': 'ویرایش',
        }
        self.assertEqual(self.client.post(reverse('edit', args=[obj.pk]), p).status_code, 302)
        obj.refresh_from_db()
        self.assertEqual(obj.area, D('55.25'))
        self.assertEqual(self.client.get(reverse('user_edit', args=[self.entry.pk])).status_code, 200)
        self.assertEqual(self.client.post('/users/', {
            'username': 'newuser',
            'role': 'viewer',
            'password': 'VeryStrong!Password93',
            'is_active': 'on',
            'factories': [self.f1.pk],
        }).status_code, 302)

@skipUnless(connection.vendor == 'postgresql', 'Requires PostgreSQL row locking; run scripts/test-postgres.sh')
class PostgreSQLConcurrencyTests(TransactionTestCase):
    reset_sequences = True

    def test_two_identical_concurrent_submissions(self):
        from concurrent.futures import ThreadPoolExecutor
        user = get_user_model().objects.create_superuser('concurrent', password='password')
        f = Factory.objects.create(name='کارخانه')
        s = Size.objects.create(width=60, length=60)
        g = Grade.objects.create(name='درجه ۱', rank=1)
        key = uuid.uuid4()
        def submit(_):
            close_old_connections()
            try:
                return create_batch(user, f, date(2026, 9, 9), [{'size': s, 'grade': g, 'area': D(100), 'notes': ''}], key)
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, range(2)))
        self.assertEqual(sorted(results), [False, True])
        self.assertEqual(Production.objects.count(), 1)
        self.assertEqual(Audit.objects.count(), 1)
