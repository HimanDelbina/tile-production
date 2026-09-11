import os
import openpyxl
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Optional
from django.core.exceptions import ValidationError
from .parser import normalize_text

from .validation import validate_sheet_limits

def read_excel_import_file(file_path: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """Read and validate Excel file for tile production import.
    
    Returns list of row dicts:
        {
            'excel_row': int,
            'raw_description': str,
            'quantity': Optional[Decimal],
            'area': Optional[Decimal],
            'formula_error': Optional[str],
            'quantity_error': Optional[str],
            'area_error': Optional[str],
            'source_errors': List[Dict[str, str]],
        }
    """
    if not os.path.exists(file_path):
        raise ValidationError("فایل اکسل مشخص‌شده یافت نشد.")
        
    try:
        # Load workbook with data_only=False to inspect formulas
        wb_raw = openpyxl.load_workbook(filename=file_path, data_only=False)
        # Load workbook with data_only=True to read cached values if needed
        wb_data = openpyxl.load_workbook(filename=file_path, data_only=True)
    except Exception as e:
        raise ValidationError(f"فایل اکسل ارسالی نامعتبر یا آسیب‌دیده است: {e}")
        
    # Select sheet
    target_sheet_name = None
    if sheet_name and sheet_name.strip():
        if sheet_name.strip() in wb_raw.sheetnames:
            target_sheet_name = sheet_name.strip()
        else:
            raise ValidationError(f"شیت با نام «{sheet_name}» در فایل اکسل یافت نشد. شیت‌های موجود: {', '.join(wb_raw.sheetnames)}")
            
    ws_raw = wb_raw[target_sheet_name] if target_sheet_name else wb_raw.active
    ws_data = wb_data[target_sheet_name] if target_sheet_name else wb_data.active
    
    # Validate sheet limits
    validate_sheet_limits(ws_raw)
    
    # 1. Map columns from header row
    header_row_raw = next(ws_raw.iter_rows(min_row=1, max_row=1, values_only=False), None)
    if not header_row_raw or all(c.value is None for c in header_row_raw):
        raise ValidationError("سطر سربرگ در فایل اکسل خالی است.")
        
    desc_idx = None
    qty_idx = None
    area_idx = None
    
    for idx, cell in enumerate(header_row_raw):
        if cell.value is None:
            continue
        header_text = normalize_text(str(cell.value)).replace(" ", "")
        if "شرح" in header_text or "کالا" in header_text:
            desc_idx = idx
        elif "جمع" in header_text or "تولید" in header_text:
            qty_idx = idx
        elif "متراژ" in header_text:
            area_idx = idx
            
    missing_cols = []
    if desc_idx is None:
        missing_cols.append("«شرح کالا»")
    if qty_idx is None:
        missing_cols.append("«جمع تولید»")
    if area_idx is None:
        missing_cols.append("«متراژ»")
        
    if missing_cols:
        raise ValidationError(f"ستون‌های ضروری زیر در فایل اکسل یافت نشدند: {', '.join(missing_cols)}")
        
    # 2. Iterate data rows
    rows_raw = list(ws_raw.iter_rows(min_row=2, values_only=False))
    rows_data = list(ws_data.iter_rows(min_row=2, values_only=False))
    
    parsed_rows = []
    
    for row_num, (r_raw, r_data) in enumerate(zip(rows_raw, rows_data), start=2):
        # Check if entire row is empty
        if all(c.value is None or str(c.value).strip() == "" for c in r_raw):
            continue
            
        cell_desc_raw = r_raw[desc_idx] if desc_idx < len(r_raw) else None
        cell_qty_raw = r_raw[qty_idx] if qty_idx < len(r_raw) else None
        cell_area_raw = r_raw[area_idx] if area_idx < len(r_raw) else None
        
        cell_desc_data = r_data[desc_idx] if desc_idx < len(r_data) else None
        cell_qty_data = r_data[qty_idx] if qty_idx < len(r_data) else None
        cell_area_data = r_data[area_idx] if area_idx < len(r_data) else None
        
        source_errors = []
        
        # Check for formula in essential cells
        formula_error = None
        for col_name, c_raw in [("شرح کالا", cell_desc_raw), ("جمع تولید", cell_qty_raw), ("متراژ", cell_area_raw)]:
            if c_raw and (c_raw.data_type == 'f' or str(c_raw.value or '').startswith('=')):
                formula_error = f"سلول {col_name} در ردیف {row_num} حاوی فرمول است؛ سلول‌های ضروری نباید فرمول باشند."
                source_errors.append({"code": "formula", "message": formula_error})
                break
                
        # Parse description
        raw_desc = str(cell_desc_data.value if cell_desc_data and cell_desc_data.value is not None else "").strip()
        
        # Parse quantity
        quantity_val = None
        qty_err = None
        if cell_qty_data and cell_qty_data.value not in (None, ""):
            try:
                norm_q = normalize_text(str(cell_qty_data.value)).replace(",", "")
                quantity_val = Decimal(norm_q)
                if quantity_val < 0:
                    qty_err = "جمع تولید نمی‌تواند منفی باشد."
                    source_errors.append({"code": "negative_quantity", "message": qty_err})
            except (InvalidOperation, ValueError):
                qty_err = "مقدار جمع تولید عددی معتبر نیست."
                source_errors.append({"code": "invalid_quantity", "message": qty_err})
                
        # Parse area
        area_val = None
        area_err = None
        if cell_area_data and cell_area_data.value not in (None, ""):
            try:
                norm_a = normalize_text(str(cell_area_data.value)).replace(",", "")
                area_val = Decimal(norm_a)
                if area_val <= 0:
                    area_err = "متراژ باید بزرگ‌تر از صفر باشد."
                    source_errors.append({"code": "invalid_area", "message": area_err})
            except (InvalidOperation, ValueError):
                area_err = "مقدار متراژ عددی معتبر نیست."
                source_errors.append({"code": "invalid_area", "message": area_err})
        else:
            area_err = "متراژ الزامی است."
            source_errors.append({"code": "missing_area", "message": area_err})
            
        parsed_rows.append({
            "excel_row": row_num,
            "raw_description": raw_desc,
            "quantity": quantity_val,
            "area": area_val,
            "formula_error": formula_error,
            "quantity_error": qty_err,
            "area_error": area_err,
            "source_errors": source_errors,
        })
        
    return parsed_rows
